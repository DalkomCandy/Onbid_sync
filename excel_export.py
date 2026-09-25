from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from onbid_client import flatten_rows, format_date

HEADER_FILL = PatternFill("solid", fgColor="173A63")
HEADER_FONT = Font(name="Malgun Gothic", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Malgun Gothic", size=10)
THIN = Border(
    left=Side(style="thin", color="D6D3CD"),
    right=Side(style="thin", color="D6D3CD"),
    top=Side(style="thin", color="D6D3CD"),
    bottom=Side(style="thin", color="D6D3CD"),
)
STATUS_FILL = {
    "수의": PatternFill("solid", fgColor="FFF3C4"),
    "유찰": PatternFill("solid", fgColor="FFE4E6"),
    "낙찰": PatternFill("solid", fgColor="D1FAE5"),
    "조회": PatternFill("solid", fgColor="E7E5E4"),
}


def _style_header(ws, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.auto_filter.ref = f"A1:{get_column_letter(col_count)}1"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _status_fill(status: str | None):
    text = status or ""
    for key, fill in STATUS_FILL.items():
        if key in text:
            return fill
    return None


def write_excel(snapshot: dict[str, Any], dest: Path) -> Path:
    wb = Workbook()

    items = wb.active
    items.title = "물건목록"
    headers = [
        "표 공고번호",
        "온비드 공고번호",
        "물건관리번호",
        "물건명",
        "진행상태",
        "회차",
        "입찰시작",
        "입찰종료",
        "최저입찰가",
        "감정평가액",
        "공고기관",
        "온비드 링크",
        "비고",
    ]
    items.append(headers)
    _style_header(items, len(headers))
    for row in flatten_rows(snapshot):
        items.append(
            [
                row["originalPbanc"],
                row["pbancMngNo"],
                row["cltrMngNo"],
                row["name"],
                row["status"],
                row["round"],
                row["bidStart"],
                row["bidEnd"],
                row["minBid"],
                row["appraisal"],
                row["org"],
                row["url"],
                row["note"],
            ]
        )
        last = items.max_row
        fill = _status_fill(row["status"])
        if fill:
            items.cell(last, 5).fill = fill
        items.cell(last, 9).number_format = "#,##0"
        items.cell(last, 10).number_format = "#,##0"

    widths = [18, 18, 20, 48, 14, 8, 12, 12, 16, 16, 22, 28, 36]
    for idx, width in enumerate(widths, start=1):
        items.column_dimensions[get_column_letter(idx)].width = width
    for row in items.iter_rows(min_row=2, max_row=items.max_row, max_col=len(headers)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    notices = wb.create_sheet("공고목록")
    notice_headers = [
        "표 공고번호",
        "온비드 공고번호",
        "조회여부",
        "공고명",
        "공고기관",
        "진행상태",
        "최신회차",
        "물건수",
        "입찰시작",
        "입찰종료",
        "온비드 링크",
        "비고",
    ]
    notices.append(notice_headers)
    _style_header(notices, len(notice_headers))
    for ann in snapshot.get("announcements") or []:
        notices.append(
            [
                ann.get("originalPbanc"),
                ann.get("currentPbanc"),
                "조회됨" if ann.get("found") else "조회불가",
                ann.get("title"),
                ann.get("org"),
                ann.get("latestStatus"),
                ann.get("latestRound"),
                ann.get("itemCount"),
                format_date(ann.get("bidStart")),
                format_date(ann.get("bidEnd")),
                ann.get("onbidUrl"),
                ann.get("note"),
            ]
        )
    for idx, width in enumerate([18, 18, 10, 52, 22, 14, 10, 10, 12, 12, 28, 48], start=1):
        notices.column_dimensions[get_column_letter(idx)].width = width
    for row in notices.iter_rows(min_row=2, max_row=notices.max_row, max_col=len(notice_headers)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = THIN

    schedule = wb.create_sheet("입찰일정")
    schedule.append(["온비드 공고번호", "공고명", "회차", "입찰시작", "입찰종료", "개찰일시", "진행상태", "입찰방식"])
    _style_header(schedule, 8)
    for ann in snapshot.get("announcements") or []:
        for rnd in ann.get("schedule") or []:
            schedule.append(
                [
                    ann.get("currentPbanc"),
                    ann.get("title"),
                    rnd.get("pbctNsq"),
                    rnd.get("pbctBgngDt"),
                    rnd.get("pbctDdlnDt"),
                    rnd.get("opbxDt"),
                    rnd.get("exctStatNm"),
                    rnd.get("bidMthodNm"),
                ]
            )
    for idx, width in enumerate([18, 52, 8, 20, 20, 20, 14, 14], start=1):
        schedule.column_dimensions[get_column_letter(idx)].width = width

    meta = wb.create_sheet("조회정보")
    meta.append(["항목", "값"])
    _style_header(meta, 2)
    meta.append(["조회시각", snapshot.get("fetchedAt")])
    meta.append(["출처", snapshot.get("source")])
    meta.append(["추적 공고", snapshot.get("announcementCount")])
    meta.append(["조회된 공고", snapshot.get("foundCount")])
    meta.append(["조회불가 공고", snapshot.get("missingCount")])
    meta.append(["물건 수(최신 회차)", snapshot.get("itemCount")])
    meta.column_dimensions["A"].width = 22
    meta.column_dimensions["B"].width = 48

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def default_excel_path(snapshot: dict[str, Any] | None = None) -> Path:
    stamp = (snapshot or {}).get("fetchedAt") or datetime.now().isoformat()
    date = str(stamp)[:10].replace("-", "")
    return Path.home() / "Downloads" / f"온비드_공매현황_{date}.xlsx"
